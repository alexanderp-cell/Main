#!/usr/bin/env python3
"""Собирает клиентскую таблицу UT PO Status из внутренней.

Внутренний файл (UT_PO_STATUS_<дата>.xlsx) — источник правды.
Предыдущая клиентская книга — то, что уже видел клиент: ручные комментарии
и строки без изменений сохраняются.
Шаблон задаёт шрифт, ширины, условное форматирование и выпадающие списки.

Пример:
  python3 build_client_po_status.py \\
    UT_PO_STATUS_13.10.2026.xlsx \\
    --previous Ready_UT_PO_STATUS_06.10.2026.xlsx \\
    --template template/client_po_status_template.xlsx
"""

from __future__ import annotations

import argparse
import json
import re
import shutil
import zipfile
from collections import defaultdict
from copy import copy
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl

SECTIONS = [
    "Ожидает аванса",
    "Требует внимания",
    "Принят в работу",
    "Отгрузка от поставщика",
    "Лид тайм",
    "В транзите",
    "Отгружен",
    "Завершен",
    "Отменен",
]

SECTION_HEADERS = {
    "TO BE SOLVED": "TO BE SOLVED",
    "TO BE PLACED": "TO BE PLACED",
    "TROUBLES": "TROUBLES",
    "MARKET": "MARKET",
    "BLUE OCEAN SOLUTIONS": "BLUE OCEAN SOLUTIONS",
    "PRE TRANSIT": "PRE TRANSIT",
    "SHIPPED": "SHIPPED",
    "FINISHED": "FINISHED",
}

ACTIVE_SOURCE_KINDS = {"trouble", "paid", "not_paid", "shipped"}
ARCHIVE_FLIP_KINDS = {"shipped", "finished", "cancelled", "refund", "warranty", "solved", "bis", "other"}

COMMENT_WHS = "На складе консолидации"
COMMENT_ENROUTE = "В пути в транзитную точку"
COMMENT_POINT = "в транзитной точке"
COMMENT_DELAY = "Задержка с оформлением документов в транзитной точке"

SHIPPED_MAX_AGE_DAYS = 60


@dataclass
class SourceRow:
    row: int
    section: str
    kind: str
    invoice: str
    pn_key: str
    pn_raw: str
    alt: object
    desc: object
    qty: object
    start: date | None
    lead: object
    deadline: date | None
    actual: date | None
    condition: object
    price: object
    total: object
    payment: object
    awb: object
    arrival: object
    comment: str
    warranty: bool
    on_hold: bool
    row_color: str | None = None
    serial: object = None
    batch_planned: date | None = None
    batch_comment: str | None = None


@dataclass
class PrevRow:
    row: int
    section: str
    values: list
    inv: str
    pn: str
    qty: object


@dataclass
class Decision:
    action: str  # keep, update, add, drop, skip
    section: str | None = None
    rule: str = ""
    comment: object = None
    comment_mode: str = "keep"  # keep, set, clear
    planned: object = None  # date | "N/A" | None
    planned_mode: str = "keep"  # keep, set, clear
    fill_dates: bool = True  # start + deadline
    lock_dates: bool = False
    actual: date | None = None
    actual_mode: str = "keep"  # keep, set
    refresh_commercial: bool = False


def norm_inv(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    if text.endswith(".0"):
        text = text[:-2]
    return text


def norm_pn(value) -> str:
    if value is None:
        return ""
    return " ".join(str(value).strip().upper().split())


def norm_qty(value):
    if value is None or value == "":
        return None
    if isinstance(value, str):
        text = value.strip().replace(" ", "").replace(",", ".")
        if not text:
            return None
        try:
            value = float(text)
        except ValueError:
            return value.strip()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, int):
        return value
    return value


def parse_stamp(value) -> datetime | None:
    """Дата из ячейки: объект Excel или текст ДД.ММ.ГГГГ, в том числе со временем."""
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    if not isinstance(value, str):
        return None
    text = value.strip()
    if not text or text.startswith("="):
        return None
    text = re.sub(r"\b(\d):", r"0\1:", text)
    for fmt in (
        "%d.%m.%Y",
        "%d.%m.%Y %H:%M",
        "%d.%m.%Y %H:%M:%S",
        "%Y-%m-%d",
        "%Y-%m-%d %H:%M",
        "%Y-%m-%d %H:%M:%S",
    ):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def as_date(value) -> date | None:
    found = parse_stamp(value)
    if found is None or found.year < 2024:
        return None
    return found.date()


def parse_client_date(value) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, str) and value.strip().upper() in {"N/A", "NA", "-"}:
        return None
    found = as_date(value)
    if found:
        return found
    if isinstance(value, str):
        match = re.match(r"(\d{1,2})\.(\d{1,2})\.(\d{4})", value.strip())
        if match:
            day, month, year = (int(part) for part in match.groups())
            try:
                return date(year, month, day)
            except ValueError:
                return None
    return None


def fmt_date(value: date | None):
    if value is None:
        return None
    return value.strftime("%d.%m.%Y")


def file_date_from_name(path: Path) -> date:
    match = re.search(r"(\d{2})\.(\d{2})\.(\d{4})", path.name)
    if not match:
        raise SystemExit(f"В имени файла нет даты ДД.ММ.ГГГГ: {path.name}")
    day, month, year = (int(part) for part in match.groups())
    return date(year, month, day)


def status_kind(value) -> str:
    text = str(value or "").upper()
    if "REFUND" in text:
        return "refund"
    if "CANCEL" in text:
        return "cancelled"
    if "WARRANTY" in text:
        return "warranty"
    if "TROUBLE" in text:
        return "trouble"
    if "SHIPPED" in text:
        return "shipped"
    if "FINISHED" in text:
        return "finished"
    if "NOT PAID" in text:
        return "not_paid"
    if re.search(r"\bPAID\b", text):
        return "paid"
    if "SOLVED" in text:
        return "solved"
    if "BIS" in text:
        return "bis"
    if not text.strip():
        return "empty"
    return "other"


def is_warranty_text(*parts) -> bool:
    blob = "\n".join(str(part or "") for part in parts).lower()
    return "гарант" in blob or "warranty" in blob or "возврат из ремонта" in blob


def is_on_hold(*parts) -> bool:
    blob = "\n".join(str(part or "") for part in parts).upper()
    return "ON HOLD HA WHS" in blob


def payment_kind(value) -> str:
    text = str(value or "").lower()
    if "частич" in text:
        return "partial"
    if "постоплат" in text:
        return "postpay"
    if "предоплат" in text:
        return "prepay"
    return "other"


def map_payment(value):
    kind = payment_kind(value)
    if kind == "partial":
        return "Частичная предоплата"
    if kind == "postpay":
        return "Постоплата"
    if kind == "prepay":
        return "Предоплата"
    if value is None or str(value).strip().upper() in {"", "N/A", "NA"}:
        return None
    return str(value).strip()


def money(value):
    if value is None or value == "":
        return None
    if isinstance(value, str):
        text = value.strip().replace(" ", "").replace(",", ".")
        if not text:
            return None
        try:
            value = float(text)
        except ValueError:
            return value.strip()
    try:
        return f"{round(float(value) + 1e-9, 2):.2f}".replace(".", ",")
    except (TypeError, ValueError):
        return value


def money_amount(value):
    if value is None or value == "":
        return None
    if isinstance(value, str):
        text = value.strip().replace(" ", "").replace(",", ".")
        if not text:
            return None
        try:
            return round(float(text), 2)
        except ValueError:
            return None
    try:
        return round(float(value), 2)
    except (TypeError, ValueError):
        return None


def order_number(comment: str):
    """Только базис поставки и номер заказа: «DDP MOW» и «P...»."""
    if not comment:
        return None
    basis = None
    po = None
    basis_re = re.compile(r"\b(DAT|DDP|DAP|EXW|FCA|CIP|CPT|FOB)\s+([A-ZА-ЯЁ]+)", re.I)
    po_re = re.compile(r"\bP(\d{4,})\b", re.I)
    for raw in str(comment).replace("\r", "").split("\n"):
        line = raw.strip()
        if not line:
            continue
        if basis is None:
            found = basis_re.search(line)
            if found:
                basis = f"{found.group(1).upper()} {found.group(2).upper()}"
        if po is None:
            found = po_re.search(line)
            if found:
                po = f"P{found.group(1)}"
    parts = [part for part in (basis, po) if part]
    return "\n".join(parts) if parts else None


def cell_fill(cell) -> str | None:
    fill = cell.fill
    if not fill or not fill.patternType or not fill.fgColor or fill.fgColor.type != "rgb":
        return None
    rgb = fill.fgColor.rgb
    if not rgb or rgb in {"00000000", "FF000000"}:
        return None
    return rgb


def normalize_batch_token(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        if value.year < 2024:
            return ""
        return value.strftime("%Y-%m-%d %H:%M") if (value.hour or value.minute) else value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = " ".join(str(value).strip().upper().split())
    if not text or text.startswith("="):
        return ""
    return text


SHIPMENT_WORDS = ("WHS", "PAID", "READY", "HOLD", "UNSOLVED", "UNCONFIRMED", "MID-OCTOBER", "MID OCTOBER", "ON THE WAY")


def is_shipment_id(token: str) -> bool:
    if not token or any(word in token for word in SHIPMENT_WORDS):
        return False
    if parse_stamp(token):
        return False
    return len(re.sub(r"\D", "", token)) >= 8


def batch_signature(src: SourceRow):
    if src.section != "PRE TRANSIT" or src.warranty or src.on_hold:
        return None
    track = normalize_batch_token(src.awb)
    arrival = normalize_batch_token(src.arrival)
    if not track and not arrival:
        return None
    if src.row_color:
        return (src.row_color, track, arrival)
    # В новом файле строки больше не красят. Общий номер отправления — та же партия.
    if is_shipment_id(track):
        return ("TRACK", track, arrival)
    return None


def assign_batches(rows: list[SourceRow], file_date: date):
    """Одинаковый цвет и общие P/Q в PRE TRANSIT — одна консолидированная партия."""
    groups = defaultdict(list)
    for src in rows:
        signature = batch_signature(src)
        if signature:
            groups[signature].append(src)
    for members in groups.values():
        if len(members) < 2:
            continue
        shared = max(pretransit_planned(member, file_date) for member in members)
        comment = COMMENT_POINT if days_between(shared, file_date) < 5 else COMMENT_ENROUTE
        for member in members:
            member.batch_planned = shared
            member.batch_comment = comment


def parse_loose_date(value, file_date: date) -> date | None:
    found = as_date(value)
    if found:
        return found
    if not isinstance(value, str):
        return None
    match = re.search(r"(\d{1,2})[./](\d{1,2})(?:[./](\d{2,4}))?", value)
    if not match:
        return None
    day = int(match.group(1))
    month = int(match.group(2))
    year_text = match.group(3)
    if year_text is None:
        year = file_date.year
    else:
        year = int(year_text)
        if year < 100:
            year += 2000
    try:
        return date(year, month, day)
    except ValueError:
        return None


def pq_blob(src: SourceRow) -> str:
    return f"{src.awb or ''}\n{src.arrival or ''}"


def has_code(blob: str, code: str) -> bool:
    return code in blob.upper()


def comment_kind(text) -> str | None:
    if not text:
        return None
    lowered = str(text).lower()
    if "задержк" in lowered and "транзит" in lowered:
        return "delay"
    if "в пути" in lowered:
        return "enroute"
    # Длинный складской комментарий тоже содержит «транзитную точку».
    if "консолидац" in lowered or "на складе" in lowered:
        return "whs"
    if "транзитн" in lowered and "точк" in lowered:
        return "point"
    return "manual"


CANONICAL_COMMENTS = {
    "whs": COMMENT_WHS,
    "enroute": COMMENT_ENROUTE,
    "point": COMMENT_POINT,
    "delay": COMMENT_DELAY,
}


def days_between(later: date, earlier: date) -> int:
    return (later - earlier).days


def shipment_planned(deadline: date | None, file_date: date, previous: date | None, same_section: bool):
    """Планируемая дата раздела «Отгрузка от поставщика»."""
    minimum = file_date + timedelta(days=7)
    forced = False
    if deadline is not None and days_between(deadline, file_date) < 10:
        computed = deadline + timedelta(days=14)
        forced = True
    elif deadline is not None:
        computed = deadline
    else:
        computed = minimum
    if computed < minimum:
        computed = minimum
    if same_section and previous and previous >= minimum and not forced:
        return previous
    if same_section and previous and previous >= minimum and forced:
        target = deadline + timedelta(days=14) if deadline else computed
        if target < minimum:
            target = minimum
        if previous >= target:
            return previous
        return target
    return computed


def accepted_planned(deadline: date | None, file_date: date, previous: date | None, same_section: bool):
    if deadline and deadline >= file_date:
        computed = deadline
    else:
        computed = file_date + timedelta(days=7)
    if same_section and previous and previous >= file_date:
        return previous
    return computed


def whs_planned(deadline: date | None, file_date: date) -> date:
    floor = file_date + timedelta(days=10)
    if deadline is None or days_between(deadline, file_date) < 7:
        shifted = (deadline + timedelta(days=10)) if deadline else floor
        return shifted if shifted > floor else floor
    return deadline


def pretransit_planned(src: SourceRow, file_date: date) -> date:
    q_date = parse_loose_date(src.arrival, file_date)
    if q_date:
        if days_between(q_date, file_date) < 2:
            return file_date + timedelta(days=3)
        return q_date
    if src.deadline is None:
        return file_date + timedelta(days=10)
    left = days_between(src.deadline, file_date)
    if left < 0:
        return file_date + timedelta(days=10)
    return src.deadline


def apply_transit_stability(computed: date, previous: date | None, same_section: bool, last_change: date | None, file_date: date) -> date:
    if not same_section or previous is None or last_change is None:
        return computed
    if days_between(file_date, last_change) < 5 and days_between(previous, file_date) >= 7:
        return previous
    return computed


def transit_base_kind(src: SourceRow, planned: date, file_date: date) -> str:
    if days_between(planned, file_date) < 5:
        return "point"
    blob = pq_blob(src).upper()
    in_pretransit = src.section == "PRE TRANSIT"
    whs = any(token in blob for token in ("BOS WHS", "HA WHS", "НА WHS")) or re.search(r"\bWHS\b", blob)
    # «On the way» и UNSOLVED сюда не попадают: их разбирает классификатор BOS раньше.
    if whs and not in_pretransit:
        return "whs"
    return "enroute"


def choose_transit_comment(base_kind: str, prev_text, planned_changed: bool, point_changes: int):
    prev_kind = comment_kind(prev_text)
    kind = base_kind
    next_point_changes = point_changes
    if prev_kind == "delay":
        kind = "delay"
    elif prev_kind == "point" and planned_changed and point_changes + 1 >= 2:
        kind = "delay"
        next_point_changes = point_changes + 1
    elif kind == "point" and prev_kind == "point" and planned_changed:
        next_point_changes = point_changes + 1
    elif kind == "point" and prev_kind != "point":
        next_point_changes = point_changes

    if kind == prev_kind and prev_text:
        return prev_text, kind, next_point_changes
    return CANONICAL_COMMENTS[kind], kind, next_point_changes


def awb_comment(src: SourceRow):
    awb = str(src.awb).strip() if src.awb else ""
    arrival = parse_stamp(src.arrival)
    if arrival and (arrival.hour or arrival.minute):
        when = arrival.strftime("%d.%m.%Y %H:%M")
    else:
        when_date = as_date(src.arrival)
        when = fmt_date(when_date) if when_date else ""
    if awb and when:
        return f"AWB {awb}\n{when}"
    if awb:
        return f"AWB {awb}"
    return None


COLUMN_SPECS = (
    ("invoice", ("номер счета", "invoice")),
    ("status", ("status", "статус")),
    ("comment", ("комментар",)),
    ("pn", ("p/n", "part number", "парт")),
    ("alt", ("alt",)),
    ("desc", ("description", "описан")),
    ("qty", ("qty", "колич")),
    ("start", ("взят в работу",)),
    ("lead", ("дней на поставку",)),
    ("deadline", ("крайняя дата",)),
    ("actual", ("фактическая дата",)),
    ("condition", ("condition", "состояние")),
    ("serial", ("серийн", "s/n", "s.n")),
    ("price", ("продажная, ед", "продажная,ед", "unit price")),
    ("total", ("продажная, итого", "итого")),
    ("payment", ("тип оплаты", "payment")),
    ("awb", ("awb",)),
    ("arrival", ("ожидаемая дата", "прихода")),
)


def header_columns(ws) -> dict[str, int]:
    """Столбцы ищем по заголовку: в новом файле добавлены Client, UOM и S/N."""
    found = {}
    for col in range(1, ws.max_column + 1):
        header = " ".join(str(ws.cell(1, col).value or "").strip().lower().split())
        if not header:
            continue
        for key, needles in COLUMN_SPECS:
            if key in found:
                continue
            if any(needle in header for needle in needles):
                found[key] = col
                break
    required = ("invoice", "status", "comment", "pn", "qty", "start", "lead", "deadline", "price", "payment", "awb", "arrival")
    missing = [name for name in required if name not in found]
    if missing:
        raise SystemExit(f"В таблице 1 не найдены столбцы: {', '.join(missing)}")
    return found


def display_code(value):
    """Число из Excel без хвоста .0: партномер, трек и серийный номер."""
    if value is None or value == "":
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int) and not isinstance(value, bool):
        return str(value)
    text = str(value).replace("\r\n", "\n").strip()
    if re.fullmatch(r"\d+\.0", text):
        return text[:-2]
    return text or None


def pure_date_value(value):
    """Ячейка целиком является датой, а не комментарием с датой внутри."""
    if parse_stamp(value) is None:
        return False
    if isinstance(value, (datetime, date)):
        return True
    text = str(value).strip()
    text = re.sub(r"\b(\d):", r"0\1:", text)
    return bool(re.fullmatch(r"\d{1,2}\.\d{1,2}\.\d{4}(?:\s+\d{2}:\d{2}(?::\d{2})?)?|\d{4}-\d{2}-\d{2}(?:\s+\d{2}:\d{2}(?::\d{2})?)?", text))


def drop_shadow_duplicates(rows: list[SourceRow]) -> list[SourceRow]:
    """Одна и та же строка в TROUBLES и ещё раз в MARKET — для клиента оставляем TROUBLES."""
    groups = defaultdict(list)
    for src in rows:
        groups[(src.invoice, src.pn_key, norm_qty(src.qty))].append(src)
    drop = set()
    for group in groups.values():
        troubles = [src for src in group if src.section == "TROUBLES" or src.kind == "trouble"]
        if not troubles:
            continue
        for src in group:
            if src not in troubles and src.kind in ACTIVE_SOURCE_KINDS:
                drop.add(id(src))
    return [src for src in rows if id(src) not in drop]


def load_source(path: Path) -> list[SourceRow]:
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb.active
    cols = header_columns(ws)
    section = "PREAMBLE"
    rows: list[SourceRow] = []

    def cell(index, key):
        column = cols.get(key)
        if not column:
            return None
        return ws.cell(index, column).value

    for index in range(2, ws.max_row + 1):
        invoice_raw = cell(index, "invoice")
        status = cell(index, "status")
        comment = cell(index, "comment")
        pn = cell(index, "pn")
        kind = status_kind(status)
        header = str(pn).strip().upper() if isinstance(pn, str) else ""
        invoice = norm_inv(invoice_raw)
        if header in SECTION_HEADERS and not invoice:
            section = SECTION_HEADERS[header]
            continue
        blank = not invoice and not norm_pn(pn) and kind in {"empty", "other"}
        if section == "PRE TRANSIT" and blank:
            section = "ARCHIVE"
            continue
        if section == "PRE TRANSIT" and kind in ARCHIVE_FLIP_KINDS:
            section = "ARCHIVE"
        if str(comment or "").strip().upper() == "SAMPLE":
            continue
        pn_text = display_code(pn)
        if not invoice or not norm_pn(pn_text) or invoice in {"-", "NONE"}:
            continue
        start = as_date(cell(index, "start"))
        lead = cell(index, "lead")
        deadline = as_date(cell(index, "deadline"))
        if deadline is None and start is not None and isinstance(lead, (int, float)):
            deadline = start + timedelta(days=int(lead))
        awb = cell(index, "awb")
        arrival = cell(index, "arrival")
        # В новом файле ожидаемую дату иногда пишут в столбец AWB, а столбец прихода оставляют пустым.
        if pure_date_value(awb) and not pure_date_value(arrival):
            arrival = awb
            awb = None
        else:
            awb = display_code(awb)
        item = SourceRow(
            row=index,
            section=section,
            kind=kind,
            invoice=invoice,
            pn_key=norm_pn(pn_text),
            pn_raw=pn_text,
            alt=cell(index, "alt"),
            desc=cell(index, "desc"),
            qty=cell(index, "qty"),
            start=start,
            lead=lead,
            deadline=deadline,
            actual=as_date(cell(index, "actual")),
            condition=cell(index, "condition"),
            price=cell(index, "price"),
            total=cell(index, "total"),
            payment=cell(index, "payment"),
            awb=awb,
            arrival=arrival,
            comment=str(comment or ""),
            warranty=is_warranty_text(comment, invoice_raw, pn_text),
            on_hold=is_on_hold(awb, arrival, comment),
            row_color=cell_fill(ws.cell(index, cols["invoice"])),
            serial=display_code(cell(index, "serial")),
        )
        rows.append(item)
    return drop_shadow_duplicates(rows)


def load_previous(path: Path) -> list[PrevRow]:
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb["Main"] if "Main" in wb.sheetnames else wb.active
    banners = []
    for index in range(1, ws.max_row + 1):
        pn = ws.cell(index, 6).value
        invoice = ws.cell(index, 4).value
        desc = ws.cell(index, 8).value
        if isinstance(pn, str) and invoice in (None, "") and desc in (None, ""):
            banners.append((index, pn.strip()))
    rows = []
    for index in range(2, ws.max_row + 1):
        invoice = ws.cell(index, 4).value
        if invoice in (None, ""):
            continue
        section = ""
        for banner_row, name in banners:
            if banner_row < index:
                section = name
        values = [ws.cell(index, col).value for col in range(1, 25)]
        rows.append(
            PrevRow(
                row=index,
                section=section,
                values=values,
                inv=norm_inv(invoice),
                pn=norm_pn(ws.cell(index, 6).value),
                qty=norm_qty(ws.cell(index, 9).value),
            )
        )
    return rows


def market_section(src: SourceRow, file_date: date) -> tuple[str, str]:
    pay = payment_kind(src.payment)
    age = days_between(file_date, src.start) if src.start else 0
    if src.kind == "paid":
        return "Отгрузка от поставщика", "MARKET: оплачен → отгрузка от поставщика"
    if pay == "partial" or pay == "prepay":
        if age < 10:
            return "Ожидает аванса", "MARKET: частичная предоплата, меньше 10 дней → ожидает аванса"
        return "Принят в работу", "MARKET: частичная предоплата, 10 дней и больше → принят в работу"
    if pay == "postpay":
        if age > 10:
            return "Отгрузка от поставщика", "MARKET: 100% постоплата, больше 10 дней → отгрузка от поставщика"
        return "Принят в работу", "MARKET: 100% постоплата, до 10 дней → принят в работу"
    if age < 10:
        return "Ожидает аванса", "MARKET: тип оплаты не распознан, заказ свежий → ожидает аванса"
    return "Принят в работу", "MARKET: тип оплаты не распознан → принят в работу"


def bos_section(src: SourceRow, file_date: date) -> tuple[str, str] | None:
    blob = pq_blob(src).upper()
    if "ON HOLD HA WHS" in blob:
        return None
    if "UNSOLVED" in blob:
        return "Требует внимания", "BOS: UNSOLVED → требует внимания"
    if re.search(r"\bPAID\b", blob) or "ON THE WAY TO WHS" in blob:
        return "Отгрузка от поставщика", "BOS: PAID или On the way to whs → отгрузка от поставщика"
    if any(token in blob for token in ("BOS WHS", "HA WHS", "НА WHS")) or re.search(r"\bWHS\b", blob):
        return "В транзите", "BOS: склад консолидации → в транзите"
    if "UNCONFIRMED YET" in blob or not str(src.awb or "").strip():
        section, _rule = market_section(replace_kind(src, "not_paid"), file_date)
        return section, f"BOS: UNCONFIRMED или пусто, далее правило MARKET → {section}"
    section, _rule = market_section(replace_kind(src, "not_paid"), file_date)
    return section, f"BOS: по условию оплаты → {section}"


def replace_kind(src: SourceRow, kind: str) -> SourceRow:
    clone = SourceRow(**{**src.__dict__})
    clone.kind = kind
    return clone


def is_recent_shipped(src: SourceRow, file_date: date) -> bool:
    arrived = as_date(src.arrival)
    if arrived is None:
        return True
    return days_between(file_date, arrived) <= SHIPPED_MAX_AGE_DAYS


def classify(src: SourceRow | None, prev: PrevRow | None, file_date: date, history: dict) -> Decision:
    if src is None:
        return Decision(action="keep", section=prev.section if prev else None, rule="В новом файле строки нет, оставлена как была")

    if src.warranty or src.kind == "warranty":
        if prev:
            return Decision(action="drop", rule="Гарантия: строка убирается из клиентской таблицы")
        return Decision(action="skip", rule="Гарантия")
    if src.on_hold or (src.section == "BLUE OCEAN SOLUTIONS" and "ON HOLD HA WHS" in pq_blob(src).upper()):
        if prev:
            return Decision(action="drop", rule="ON HOLD HA WHS: строка не показывается клиенту")
        return Decision(action="skip", rule="ON HOLD HA WHS")
    if src.kind in {"cancelled", "refund"}:
        if prev and prev.section not in {"Завершен", "Отменен"}:
            return Decision(
                action="update",
                section="Отменен",
                rule="Refund/Cancelled: ранее показанная строка перенесена в «Отменен»",
                comment_mode="keep",
                planned="N/A",
                planned_mode="set",
                fill_dates=True,
                refresh_commercial=False,
            )
        if prev:
            return Decision(
                action="keep",
                section=prev.section,
                rule="Refund/Cancelled у уже закрытой строки не учитывается",
            )
        return Decision(action="skip", rule="Refund/Cancelled не добавляются")

    if src.kind == "shipped":
        if prev and prev.section == "Завершен":
            return Decision(
                action="keep",
                section="Завершен",
                rule="Строка уже завершена, статус SHIPPED её не открывает",
            )
        if prev is None and not is_recent_shipped(src, file_date):
            return Decision(action="skip", rule="Старый SHIPPED вне клиентской таблицы")
        arrived = as_date(src.arrival)
        return Decision(
            action="update" if prev else "add",
            section="Отгружен",
            rule="SHIPPED → отгружен",
            comment=awb_comment(src),
            comment_mode="set",
            planned_mode="keep",
            actual=arrived,
            actual_mode="set" if arrived else "keep",
            refresh_commercial=True,
            fill_dates=True,
        )

    if src.kind == "finished":
        if prev is None:
            return Decision(action="skip", rule="Архивный FINISHED не добавляется")
        arrived = as_date(src.arrival) or src.actual
        comment = awb_comment(src) if not (prev.values[0]) else None
        return Decision(
            action="update",
            section="Завершен",
            rule="FINISHED → завершен",
            comment=comment,
            comment_mode="set" if comment else "keep",
            planned_mode="keep",
            actual=arrived,
            actual_mode="set" if arrived and not prev.values[13] else "keep",
            refresh_commercial=False,
            fill_dates=True,
        )

    if src.kind == "bis":
        return Decision(action="skip", rule="Дубль BIS")

    if src.section == "TROUBLES" or src.kind == "trouble":
        return Decision(
            action="update" if prev else "add",
            section="Требует внимания",
            rule="TROUBLES → требует внимания, комментарий пустой",
            comment_mode="clear" if not (prev and prev.section == "Требует внимания") else "keep",
            planned_mode="clear",
            fill_dates=True,
            refresh_commercial=True,
        )

    if "mid-october" in f"{src.awb} {src.arrival} {src.comment}".lower() or "mid october" in f"{src.awb} {src.arrival}".lower():
        return Decision(
            action="update" if prev else "add",
            section="Лид тайм",
            rule="PRE TRANSIT: mid-october → лид тайм",
            comment_mode="keep",
            planned="N/A",
            planned_mode="set" if not (prev and prev.values[1]) else "keep",
            fill_dates=True,
            refresh_commercial=True,
        )

    target = None
    rule = ""
    if src.section == "MARKET":
        target, rule = market_section(src, file_date)
    elif src.section == "BLUE OCEAN SOLUTIONS":
        found = bos_section(src, file_date)
        if found is None:
            if prev:
                return Decision(action="drop", rule="ON HOLD HA WHS")
            return Decision(action="skip", rule="ON HOLD HA WHS")
        target, rule = found
    elif src.section == "PRE TRANSIT":
        target, rule = "В транзите", "PRE TRANSIT → в транзите"
    elif src.kind in {"paid", "not_paid"}:
        target, rule = market_section(src, file_date)
        rule = "Строка вне разделов, применено правило MARKET. " + rule
    else:
        if prev:
            return Decision(action="keep", section=prev.section, rule="Статус не разобран, строка оставлена")
        return Decision(action="skip", rule="Статус не разобран")

    decision = Decision(
        action="update" if prev else "add",
        section=target,
        rule=rule,
        refresh_commercial=True,
        fill_dates=target != "Ожидает аванса",
    )
    if target == "Ожидает аванса":
        decision.planned_mode = "clear"
        decision.comment_mode = "keep" if prev and prev.section == "Ожидает аванса" else "clear"
    elif target == "Требует внимания":
        decision.planned_mode = "clear"
        decision.comment_mode = "keep" if prev and prev.section == "Требует внимания" else "clear"
        decision.fill_dates = True
    elif target == "Лид тайм":
        decision.planned = "N/A"
        decision.planned_mode = "keep" if prev and prev.section == "Лид тайм" and prev.values[1] else "set"
        decision.comment_mode = "keep"
    elif target == "В транзите":
        decision.planned_mode = "compute-transit"
        decision.comment_mode = "compute-transit"
    elif target == "Отгрузка от поставщика":
        decision.planned_mode = "compute-shipment"
        decision.comment_mode = "keep"
    elif target == "Принят в работу":
        decision.planned_mode = "compute-accepted"
        decision.comment_mode = "keep"
    return decision


def history_key(inv: str, pn: str, qty) -> str:
    return f"{inv}|{pn}|{norm_qty(qty)}"


def history_has_date_locks(history: dict) -> bool:
    return any(bool(row.get("dates_locked")) for row in (history.get("rows") or {}).values())


def dates_locked_in_history(history: dict, inv: str, pn: str, qty) -> bool:
    rows = history.get("rows") or {}
    if rows.get(history_key(inv, pn, qty), {}).get("dates_locked"):
        return True
    prefix = f"{inv}|{pn}|"
    return any(key.startswith(prefix) and row.get("dates_locked") for key, row in rows.items())


def should_lock_dates(prev: PrevRow | None, history: dict) -> bool:
    """Дату, которая уже стоит в образце, не затираем: в таблице 1 заказ часто взят раньше оплаты.
    Пустую ячейку это не защищает — её заполняем из таблицы 1.
    """
    if prev is None:
        return False
    if not history_has_date_locks(history):
        return True
    return dates_locked_in_history(history, prev.inv, prev.pn, prev.qty)


def cell_blank(value) -> bool:
    return value in (None, "")


def apply_known_dates(values: list, prev: PrevRow | None, src: SourceRow | None, section: str | None, lock_dates: bool):
    """Дата взятия в работу и крайняя дата не остаются пустыми.
    Есть дата в образце — оставляем её. Пусто — берём из таблицы 1.
    В «Ожидает аванса» даты не ставим: только срок поставки в днях.
    """
    if section == "Ожидает аванса":
        values[10] = None
        values[12] = None
        values[13] = None
        return
    if lock_dates and prev:
        start = prev.values[10]
        deadline = prev.values[12]
    elif src:
        start = fmt_date(src.start) if src.start else (prev.values[10] if prev else None)
        deadline = fmt_date(src.deadline) if src.deadline else (prev.values[12] if prev else None)
    elif prev:
        start, deadline = prev.values[10], prev.values[12]
    else:
        start = deadline = None
    if cell_blank(start) and src and src.start:
        start = fmt_date(src.start)
    if cell_blank(deadline) and src and src.deadline:
        deadline = fmt_date(src.deadline)
    values[10] = start or None
    values[12] = deadline or None
    if section in {"Отгружен", "Завершен"} and cell_blank(values[13]) and src and src.actual:
        values[13] = fmt_date(src.actual)


def polish_client_row(values: list) -> list:
    if values[4]:
        values[4] = order_number(values[4])
    if values[2] == "Требует внимания":
        values[1] = "N/A"
    return values


def load_history(path: Path) -> dict:
    if not path.exists():
        return {"rows": {}}
    return json.loads(path.read_text(encoding="utf-8"))


def style_snapshot(cell):
    return {
        "font": copy(cell.font),
        "fill": copy(cell.fill),
        "border": copy(cell.border),
        "alignment": copy(cell.alignment),
        "number_format": cell.number_format,
    }


def apply_style(cell, snap):
    cell.font = copy(snap["font"])
    cell.fill = copy(snap["fill"])
    cell.border = copy(snap["border"])
    cell.alignment = copy(snap["alignment"])
    cell.number_format = snap["number_format"]


def build_values(prev: PrevRow | None, src: SourceRow | None, decision: Decision, file_date: date, history: dict):
    """Возвращает values, метаданные для истории и список человекочитаемых правок."""
    if decision.action == "keep" or (src is None and prev is not None):
        values = list(prev.values)
        apply_known_dates(values, prev, src, decision.section or prev.section, decision.lock_dates)
        return values, {}, []

    if (
        prev
        and decision.section in {"Завершен", "Отменен"}
        and prev.section == decision.section
        and not decision.refresh_commercial
    ):
        values = list(prev.values)
        apply_known_dates(values, prev, src, decision.section, decision.lock_dates)
        return values, {}, []

    if prev and decision.section in {"Завершен", "Отменен"} and not decision.refresh_commercial:
        values = list(prev.values)
        values[2] = decision.section
        if decision.planned_mode == "set":
            values[1] = decision.planned
        if decision.comment_mode == "set" and decision.comment and not values[0]:
            values[0] = decision.comment
        if decision.actual_mode == "set" and decision.actual and not values[13]:
            values[13] = fmt_date(decision.actual)
        apply_known_dates(values, prev, src, decision.section, decision.lock_dates)
        return values, {}, []

    values = [None] * 24
    if prev and not decision.refresh_commercial:
        values = list(prev.values)
    if src and decision.refresh_commercial:
        values[3] = coerce_invoice(src.invoice)
        parsed_order = order_number(src.comment)
        values[4] = parsed_order or (prev.values[4] if prev else None)
        values[5] = src.pn_raw
        values[6] = src.alt
        values[7] = src.desc
        values[8] = norm_qty(src.qty)
        values[9] = "EA"
        values[11] = norm_qty(src.lead) if src.lead is not None else (prev.values[11] if prev else None)
        values[14] = src.condition
        new_price = money(src.price)
        new_total = money(src.total)
        if prev and money_amount(prev.values[16]) == money_amount(src.price) and isinstance(prev.values[16], str):
            values[16] = prev.values[16]
        else:
            values[16] = new_price
        if prev and money_amount(prev.values[17]) == money_amount(src.total) and isinstance(prev.values[17], str):
            values[17] = prev.values[17]
        else:
            values[17] = new_total
        values[18] = map_payment(src.payment)
        if src.serial not in (None, ""):
            values[15] = src.serial
        elif prev:
            values[15] = prev.values[15]
        if prev:
            for col in range(19, 24):
                values[col] = prev.values[col]
    elif prev:
        values = list(prev.values)

    values[2] = decision.section
    hist_meta = {}
    key = history_key(
        norm_inv(values[3]),
        norm_pn(values[5]),
        values[8],
    )
    hist = history.get("rows", {}).get(key, {})
    last_change = parse_client_date(hist.get("last_planned_change")) if hist else None
    point_changes = int(hist.get("point_date_changes") or 0) if hist else 0
    prev_comment = prev.values[0] if prev else None
    prev_planned = parse_client_date(prev.values[1]) if prev else None
    same_section = bool(prev and prev.section == decision.section)

    if decision.comment_mode == "clear":
        if not (prev and prev.section == decision.section and prev.values[0]):
            values[0] = None
        else:
            values[0] = prev.values[0]
    elif decision.comment_mode == "set":
        values[0] = decision.comment
    elif decision.comment_mode == "keep":
        values[0] = prev.values[0] if prev else None

    if decision.fill_dates:
        apply_known_dates(values, prev, src, decision.section, decision.lock_dates)

    if decision.planned_mode == "clear":
        values[1] = None
    elif decision.planned_mode == "set":
        values[1] = decision.planned
    elif decision.planned_mode == "keep":
        values[1] = prev.values[1] if prev else None
        if decision.section == "Отгружен" and values[1] in (None, "") and src and src.deadline:
            values[1] = fmt_date(src.deadline)
    elif decision.planned_mode == "compute-shipment" and src:
        planned = shipment_planned(src.deadline, file_date, prev_planned, same_section)
        values[1] = fmt_date(planned)
    elif decision.planned_mode == "compute-accepted" and src:
        planned = accepted_planned(src.deadline, file_date, prev_planned, same_section)
        values[1] = fmt_date(planned)
    elif decision.planned_mode == "compute-transit" and src:
        if src.batch_planned:
            planned = src.batch_planned
            values[1] = fmt_date(planned)
            values[0] = src.batch_comment
            changed = prev_planned != planned
            hist_meta = {
                "comment_kind": comment_kind(src.batch_comment) or "",
                "point_date_changes": point_changes,
                "planned_changed": changed,
            }
        else:
            if src.section == "PRE TRANSIT":
                computed = pretransit_planned(src, file_date)
            else:
                computed = whs_planned(src.deadline, file_date)
            planned = apply_transit_stability(computed, prev_planned, same_section, last_change, file_date)
            values[1] = fmt_date(planned)
            base = transit_base_kind(src, planned, file_date)
            if not hist and comment_kind(prev_comment) == "delay":
                point_changes = max(point_changes, 2)
            changed = prev_planned != planned
            comment, kind, point_changes = choose_transit_comment(base, prev_comment, changed, point_changes)
            values[0] = comment
            hist_meta = {
                "comment_kind": kind,
                "point_date_changes": point_changes,
                "planned_changed": changed,
            }

    if decision.actual_mode == "set" and decision.actual:
        values[13] = fmt_date(decision.actual)
    elif decision.section == "Ожидает аванса":
        values[13] = None

    if decision.section == "Ожидает аванса":
        values[1] = None
        values[10] = None
        values[12] = None
        values[13] = None

    if decision.section == "Требует внимания":
        values[1] = "N/A"

    planned_text = values[1] if isinstance(values[1], str) else fmt_date(parse_client_date(values[1]))
    if hist_meta.get("planned_changed"):
        hist_meta["last_planned_change"] = file_date.isoformat()
    elif hist.get("last_planned_change"):
        hist_meta["last_planned_change"] = hist.get("last_planned_change")
    elif prev_planned and fmt_date(prev_planned) != planned_text:
        hist_meta["last_planned_change"] = file_date.isoformat()
    hist_meta.setdefault("comment_kind", comment_kind(values[0]) or "")
    hist_meta.setdefault("point_date_changes", point_changes)
    hist_meta["section"] = decision.section
    hist_meta["planned"] = planned_text or ""
    hist_meta["dates_locked"] = bool(decision.lock_dates)
    return values, hist_meta, []


def coerce_invoice(invoice: str):
    if invoice.isdigit() and len(invoice) <= 15:
        return int(invoice)
    return invoice


def norm_text(value) -> str:
    if value is None:
        return ""
    return str(value).replace("\r\n", "\n").strip()


def comparable(section: str, values: list) -> tuple:
    return (
        section,
        norm_text(values[0]),
        norm_text(values[1]),
        norm_text(values[2]),
        norm_inv(values[3]),
        norm_text(values[4]),
        norm_pn(values[5]),
        norm_text(values[6]),
        norm_text(values[7]),
        norm_qty(values[8]),
        norm_text(values[10]),
        norm_qty(values[11]),
        norm_text(values[12]),
        norm_text(values[13]),
        norm_text(values[14]),
        money_amount(values[16]),
        money_amount(values[17]),
        norm_text(values[18]),
    )


def describe_change(prev: PrevRow | None, section: str, values: list, rule: str) -> str | None:
    ident = f"{norm_inv(values[3])} / {values[5]}"
    if prev is None:
        return f"ДОБАВЛЕНО  {ident}  → {section}. {rule}"
    before = comparable(prev.section, prev.values)
    after = comparable(section, values)
    if before == after:
        return None
    bits = []
    if prev.section != section:
        bits.append(f"{prev.section} → {section}")
    else:
        bits.append(section)
    if norm_text(prev.values[0]) != norm_text(values[0]):
        old = norm_text(prev.values[0]) or "—"
        new = norm_text(values[0]) or "—"
        bits.append(f"комментарий: «{old}» → «{new}»")
    if norm_text(prev.values[1]) != norm_text(values[1]):
        bits.append(f"план: {norm_text(prev.values[1]) or '—'} → {norm_text(values[1]) or '—'}")
    if norm_text(prev.values[13]) != norm_text(values[13]):
        bits.append(f"факт: {norm_text(prev.values[13]) or '—'} → {norm_text(values[13]) or '—'}")
    if norm_qty(prev.values[8]) != norm_qty(values[8]):
        bits.append(f"кол-во: {prev.values[8]} → {values[8]}")
    if money_amount(prev.values[16]) != money_amount(values[16]):
        bits.append(f"цена: {prev.values[16]} → {values[16]}")
    if norm_text(prev.values[4]) != norm_text(values[4]):
        bits.append(f"заказ: «{norm_text(prev.values[4]) or '—'}» → «{norm_text(values[4]) or '—'}»")
    if norm_text(prev.values[10]) != norm_text(values[10]) or norm_text(prev.values[12]) != norm_text(values[12]):
        bits.append(
            f"даты клиента: {norm_text(prev.values[10]) or '—'} / {norm_text(prev.values[12]) or '—'} "
            f"→ {norm_text(values[10]) or '—'} / {norm_text(values[12]) or '—'}"
        )
    return f"ИЗМЕНЕНО  {ident}  {'; '.join(bits)}. {rule}"


def match_sources(previous: list[PrevRow], source_rows: list[SourceRow]):
    pools = defaultdict(list)
    for src in source_rows:
        pools[(src.invoice, src.pn_key)].append(src)
    used = set()

    def score(prev: PrevRow, src: SourceRow) -> int:
        points = 0
        if norm_qty(src.qty) == prev.qty:
            points += 100
        if prev.section == "Завершен" and src.kind == "finished":
            points += 40
        if prev.section == "Отгружен" and src.kind == "shipped":
            points += 40
        if prev.section == "Отменен" and src.kind in {"cancelled", "refund"}:
            points += 40
        if prev.section not in {"Завершен", "Отменен"} and src.kind in ACTIVE_SOURCE_KINDS | {"not_paid", "paid", "trouble"}:
            points += 25
        if src.kind == "finished" and prev.section not in {"Завершен", "Отменен", "Отгружен"}:
            points -= 15
        if src.kind == "bis":
            points -= 30
        if src.warranty:
            points += 5
        return points

    matched = {}
    for prev in previous:
        candidates = [src for src in pools.get((prev.inv, prev.pn), []) if id(src) not in used]
        if not candidates:
            matched[id(prev)] = None
            continue
        best = max(candidates, key=lambda src: score(prev, src))
        used.add(id(best))
        matched[id(prev)] = best
    unmatched = [src for src in source_rows if id(src) not in used]
    return matched, unmatched


def publishable_new(src: SourceRow, file_date: date) -> bool:
    if src.warranty or src.on_hold or src.kind in {"cancelled", "refund", "warranty", "finished", "bis", "solved", "empty"}:
        return False
    if src.kind == "shipped" and not is_recent_shipped(src, file_date):
        return False
    if src.section in {"PREAMBLE", "TO BE SOLVED", "TO BE PLACED"} and src.kind not in ACTIVE_SOURCE_KINDS:
        return False
    return src.kind in ACTIVE_SOURCE_KINDS or src.section in {
        "TROUBLES",
        "MARKET",
        "BLUE OCEAN SOLUTIONS",
        "PRE TRANSIT",
        "ARCHIVE",
    }


@dataclass
class BuiltRow:
    section: str
    values: list
    group: int
    order: int
    change: str | None = None
    rule: str = ""
    key: str = ""
    hist: dict = field(default_factory=dict)


def assemble(previous, source_rows, file_date, history):
    matched, unmatched = match_sources(previous, source_rows)
    buckets = {name: [] for name in SECTIONS}
    changes = []
    dropped = []
    notes = []
    order = 0

    for prev in previous:
        src = matched[id(prev)]
        decision = classify(src, prev, file_date, history)
        if decision.action == "keep" and src and prev.section == "Завершен" and src.kind in {"cancelled", "refund", "shipped"}:
            label = {"cancelled": "Cancelled", "refund": "Refund", "shipped": "SHIPPED"}[src.kind]
            notes.append(
                f"ОСТАВЛЕНО  {prev.inv} / {prev.values[5]}  в «Завершен». "
                f"В источнике сейчас {label}, закрытый заказ не переоткрываю."
            )
        if decision.action in {"skip"}:
            continue
        if decision.action == "drop":
            dropped.append(f"УБРАНО    {prev.inv} / {prev.values[5]}  из «{prev.section}». {decision.rule}")
            continue
        if decision.action == "keep" or decision.section is None:
            decision.section = prev.section if decision.section is None else decision.section
        if decision.section not in buckets:
            decision.section = prev.section
        decision.lock_dates = should_lock_dates(prev, history)
        if src and src.batch_planned and decision.section == "В транзите":
            decision.rule += f" Консолидированная партия, общая дата {fmt_date(src.batch_planned)}."
        values, hist, _unused = build_values(prev, src, decision, file_date, history)
        values[2] = decision.section
        polish_client_row(values)
        hist["dates_locked"] = bool(decision.lock_dates)
        change = describe_change(prev, decision.section, values, decision.rule)
        group = 0 if prev.section == decision.section else 1
        row = BuiltRow(
            section=decision.section,
            values=values,
            group=group,
            order=order,
            change=change,
            rule=decision.rule,
            key=history_key(norm_inv(values[3]), norm_pn(values[5]), values[8]),
            hist=hist,
        )
        buckets[decision.section].append(row)
        if change:
            changes.append(change)
        order += 1

    for src in unmatched:
        if not publishable_new(src, file_date):
            continue
        decision = classify(src, None, file_date, history)
        if decision.action == "skip" or not decision.section:
            continue
        decision.lock_dates = False
        if src.batch_planned and decision.section == "В транзите":
            decision.rule += f" Консолидированная партия, общая дата {fmt_date(src.batch_planned)}."
        values, hist, _unused = build_values(None, src, decision, file_date, history)
        values[2] = decision.section
        polish_client_row(values)
        hist["dates_locked"] = False
        change = describe_change(None, decision.section, values, decision.rule)
        row = BuiltRow(
            section=decision.section,
            values=values,
            group=2,
            order=order,
            change=change,
            rule=decision.rule,
            key=history_key(norm_inv(values[3]), norm_pn(values[5]), values[8]),
            hist=hist,
        )
        buckets[decision.section].append(row)
        if change:
            changes.append(change)
        order += 1

    ordered = []
    for name in SECTIONS:
        rows = sorted(buckets[name], key=lambda item: (item.group, item.order))
        ordered.append((name, rows))
    return ordered, changes, dropped, notes


def write_workbook(template: Path, output: Path, ordered, file_date: date):
    output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(template, output)
    wb = openpyxl.load_workbook(output)
    ws = wb["Main"] if "Main" in wb.sheetnames else wb.active
    data_style = [style_snapshot(ws.cell(9, col)) for col in range(1, 25)]
    banner_style = [style_snapshot(ws.cell(3, col)) for col in range(1, 25)]
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

    row_index = 2
    data_rows = []
    for section, rows in ordered:
        for col in range(1, 25):
            cell = ws.cell(row_index, col, section if col == 6 else None)
            apply_style(cell, banner_style[col - 1])
        ws.row_dimensions[row_index].height = 15
        row_index += 1
        for built in rows:
            for col in range(1, 25):
                cell = ws.cell(row_index, col, built.values[col - 1])
                apply_style(cell, data_style[col - 1])
            ws.row_dimensions[row_index].height = 28.5
            data_rows.append(row_index)
            row_index += 1

    last_row = max(row_index - 1, 2)
    ws.auto_filter.ref = f"A1:X{last_row}"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:X{last_row}"
    retarget_conditional_formatting(ws, last_row)
    wb.save(output)
    patch_dropdowns(output, last_row, data_rows)
    return last_row


def retarget_conditional_formatting(ws, last_row: int):
    rules = ws.conditional_formatting._cf_rules
    collected = []
    for sqref, cf_rules in list(rules.items()):
        if str(sqref) == "F225":
            continue
        kept = [rule for rule in cf_rules if "UNSOLVED" not in " ".join(rule.formula or [])]
        if kept:
            collected.append(kept)
    rules.clear()
    target = f"A1:X{last_row}"
    for cf_rules in collected:
        for rule in cf_rules:
            ws.conditional_formatting.add(target, rule)


def patch_dropdowns(path: Path, last_row: int, data_rows: list[int]):
    """openpyxl выбрасывает x14-валидацию. Возвращаем списки статуса и типа оплаты."""
    status_ref = compact_sqref("C", data_rows)
    pay_ref = compact_sqref("S", data_rows)
    block = (
        "<extLst><ext uri=\"{CCE6A557-97BC-4b89-ADB6-D9C93CAAB3DF}\" "
        "xmlns:x14=\"http://schemas.microsoft.com/office/spreadsheetml/2009/9/main\">"
        "<x14:dataValidations count=\"2\" xmlns:xm=\"http://schemas.microsoft.com/office/excel/2006/main\">"
        "<x14:dataValidation type=\"list\" allowBlank=\"1\" showErrorMessage=\"1\" "
        "xr:uid=\"{00000000-0002-0000-0000-000000000000}\">"
        f"<x14:formula1><xm:f>Dropdown!$A$2:$A$34</xm:f></x14:formula1><xm:sqref>{status_ref}</xm:sqref>"
        "</x14:dataValidation>"
        "<x14:dataValidation type=\"list\" allowBlank=\"1\" showErrorMessage=\"1\" "
        "xr:uid=\"{00000000-0002-0000-0000-000001000000}\">"
        f"<x14:formula1><xm:f>Dropdown!$B$2:$B$4</xm:f></x14:formula1><xm:sqref>{pay_ref}</xm:sqref>"
        "</x14:dataValidation></x14:dataValidations></ext></extLst>"
    )
    tmp = path.with_suffix(".tmp.xlsx")
    with zipfile.ZipFile(path, "r") as zin, zipfile.ZipFile(tmp, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                xml = data.decode("utf-8")
                if "xmlns:xr=" not in xml:
                    xml = xml.replace(
                        "<worksheet ",
                        '<worksheet xmlns:xr="http://schemas.microsoft.com/office/spreadsheetml/2014/revision" ',
                        1,
                    )
                xml = re.sub(r"<extLst>.*</extLst>", "", xml, count=1, flags=re.S)
                xml = xml.replace("</worksheet>", block + "</worksheet>")
                data = xml.encode("utf-8")
            zout.writestr(item, data)
    tmp.replace(path)


def compact_sqref(column: str, rows: list[int]) -> str:
    if not rows:
        return f"{column}2"
    ranges = []
    start = prev = rows[0]
    for current in rows[1:]:
        if current == prev + 1:
            prev = current
            continue
        ranges.append((start, prev))
        start = prev = current
    ranges.append((start, prev))
    parts = []
    for begin, end in ranges:
        if begin == end:
            parts.append(f"{column}{begin}")
        else:
            parts.append(f"{column}{begin}:{column}{end}")
    return " ".join(parts)


def save_history(path: Path, ordered, file_date: date, source_name: str):
    rows = {}
    for _section, items in ordered:
        for item in items:
            if not item.key:
                continue
            payload = {
                "section": item.section,
                "planned": item.hist.get("planned") or norm_text(item.values[1]),
                "comment_kind": item.hist.get("comment_kind") or "",
                "point_date_changes": item.hist.get("point_date_changes") or 0,
                "last_planned_change": item.hist.get("last_planned_change") or "",
                "dates_locked": bool(item.hist.get("dates_locked")),
            }
            rows[item.key] = payload
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(
            {"file_date": file_date.isoformat(), "source": source_name, "rows": rows},
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )


def write_changes(path: Path, file_date: date, ordered, changes, dropped, notes):
    counts = {name: len(rows) for name, rows in ordered}
    lines = [
        f"Клиентская таблица на {file_date.strftime('%d.%m.%Y')}",
        "",
        "Строк по разделам:",
    ]
    for name in SECTIONS:
        lines.append(f"  {name}: {counts.get(name, 0)}")
    lines.append("")
    lines.append(f"Изменений и новых строк: {len(changes)}")
    lines.append(f"Убрано: {len(dropped)}")
    lines.append(f"Закрытые заказы, которые не переоткрывал: {len(notes)}")
    lines.append("")
    if notes:
        lines.append("Закрытые заказы оставлены как были:")
        lines.extend(notes)
        lines.append("")
    if dropped:
        lines.append("Убранные строки:")
        lines.extend(dropped)
        lines.append("")
    if changes:
        lines.append("Отличия от прошлой клиентской таблицы:")
        lines.extend(changes)
    else:
        lines.append("Отличий от прошлой клиентской таблицы нет.")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def self_check(ordered, output: Path, file_date: date):
    index = {}
    for section, rows in ordered:
        for item in rows:
            index[(norm_inv(item.values[3]), norm_pn(item.values[5]))] = (section, item.values)
    errors = []
    if file_date == date(2026, 10, 13):
        expects = {
            ("7102610555", "J816P086A"): "Требует внимания",
            ("8102612465", "KT-1257"): "Отгрузка от поставщика",
            ("1102611295", "622-8001-350"): "В транзите",
            ("21012612372", "5100S151"): "Лид тайм",
            ("17092615215", "1150002-1"): "Отгружен",
            ("19082612405", "102LA2AG"): "Завершен",
            ("22072610315", "BACC2C4C06720EG"): "Отменен",
            ("7102613425", "2116M23P01"): "Ожидает аванса",
            ("22072613105", "3291214-2"): "Принят в работу",
            ("29092616005", "BUL216"): "Требует внимания",
            ("2309268425", "3-906S383-70"): "Отгрузка от поставщика",
            ("5102611505", "46-3145-1"): "В транзите",
            ("16092613402", "BACR15FT6AD12C"): "В транзите",
            ("23092613105", "10009255"): "В транзите",
            ("30092610505", "GR501BB4-4"): "Отгрузка от поставщика",
            ("25092613455", "121282-25"): "В транзите",
        }
        for key, section in expects.items():
            found = index.get(key)
            if not found:
                errors.append(f"нет строки {key}")
            elif found[0] != section:
                errors.append(f"{key}: ожидался раздел {section}, получен {found[0]}")
    if ("", "622-6263-003") in {(norm_pn(v[5]),) for _s, rows in ordered for v in [r.values for r in rows]}:
        errors.append("ON HOLD попал в таблицу")
    for _section, rows in ordered:
        for item in rows:
            if is_warranty_text(item.values[0], item.values[4], item.values[5]):
                errors.append(f"гарантия в выдаче: {item.values[3]} {item.values[5]}")
            if item.section == "Требует внимания" and item.values[1] != "N/A":
                errors.append(f"у «требует внимания» план должен быть N/A: {item.values[5]} = {item.values[1]!r}")
            order_text = norm_text(item.values[4])
            if order_text and any(token in order_text.lower() for token in ("предоплат", "доплат", "%", "строк")):
                errors.append(f"в номере заказа лишний текст: {item.values[5]} {order_text!r}")
            if item.section == "Ожидает аванса" and any(not cell_blank(item.values[idx]) for idx in (1, 10, 12, 13)):
                errors.append(f"в «ожидает аванса» заполнена дата: {item.values[5]}")
            if item.section not in {"Отменен", "Ожидает аванса"} and (cell_blank(item.values[10]) or cell_blank(item.values[12])):
                errors.append(f"пустая дата взятия или крайняя дата: {item.section} {item.values[3]} {item.values[5]}")
    with zipfile.ZipFile(output) as book:
        xml = book.read("xl/worksheets/sheet1.xml").decode("utf-8")
        dropdown = book.read("xl/worksheets/sheet2.xml").decode("utf-8")
    if "x14:dataValidations" not in xml:
        errors.append("пропал выпадающий список")
    if "Dropdown!$A$2:$A$34" not in xml or "Dropdown!$B$2:$B$4" not in xml:
        errors.append("формула выпадающего списка изменилась")
    if 'formula>$C1="Требует внимания"' not in xml and '$C1="Требует внимания"' not in xml:
        errors.append("пропало условное форматирование статуса")
    if "Ожидает аванса" not in dropdown or "Частичная предоплата" not in dropdown:
        errors.append("лист Dropdown поврежден")
    if index.get(("24072409372", "4051601-938")):
        errors.append("возврат из ремонта остался в таблице")
    kept = index.get(("2062611445", "002A0096-1"))
    if not kept or kept[1][10] != "10.06.2026" or kept[1][12] != "10.07.2026":
        errors.append(f"даты образца 002A0096-1 изменились: {None if not kept else (kept[1][10], kept[1][12])}")
    closed = index.get(("26032611272", "R287P04"))
    if not closed or closed[0] != "Завершен" or closed[1][10] != "26.03.2026" or closed[1][12] != "25.04.2026":
        errors.append(f"пустые даты R287P04 не взяты из таблицы 1: {None if not closed else (closed[0], closed[1][10], closed[1][12])}")
    advance = index.get(("7102613425", "2116M23P01"))
    if advance and advance[0] == "Ожидает аванса" and any(not cell_blank(advance[1][idx]) for idx in (1, 10, 12, 13)):
        errors.append(f"в авансе 2116M23P01 стоят даты: {(advance[1][1], advance[1][10], advance[1][12], advance[1][13])}")
    moved = index.get(("22072613105", "3291214-2"))
    if not moved or moved[1][10] != "22.07.2026" or moved[1][12] != "16.08.2026":
        errors.append(f"пустые даты 3291214-2 не взяты из таблицы 1: {None if not moved else (moved[1][10], moved[1][12])}")
    kept_doubt = index.get(("27072610055", "69-37325-135"))
    if not kept_doubt or kept_doubt[1][10] != "15.08.2026" or kept_doubt[1][12] != "15.09.2026":
        errors.append(f"дата образца 69-37325-135 затёрта таблицей 1: {None if not kept_doubt else (kept_doubt[1][10], kept_doubt[1][12])}")
    fresh = index.get(("7102610555", "J816P086A"))
    if not fresh or fresh[1][10] != "07.10.2026" or fresh[1][12] != "06.12.2026":
        errors.append(f"новый заказ взял не те даты: {None if not fresh else (fresh[1][10], fresh[1][12])}")
    order_cell = index.get(("7102613425", "2116M23P01"))
    if not order_cell or order_cell[1][4] != "DDP MOW\nP20657726":
        errors.append(f"номер заказа: {None if not order_cell else order_cell[1][4]!r}")
    batch = [index.get((inv, pn)) for inv, pn in (
        ("2708269485", "580-285-152"),
        ("23092613105", "10009255"),
        ("5102611165", "6364-00"),
        ("17092612085", "ST3086-013WE"),
    )]
    if any(item is None for item in batch):
        errors.append("в консолидированной партии не хватает строк")
    else:
        plans = {item[1][1] for item in batch}
        comments = {item[1][0] for item in batch}
        if len(plans) != 1 or len(comments) != 1:
            errors.append(f"партия разъехалась: план {plans}, комментарий {comments}")
    if errors:
        raise SystemExit("Проверка не пройдена:\n- " + "\n- ".join(errors))


def main():
    parser = argparse.ArgumentParser(description="Обновить клиентскую таблицу UT PO Status")
    parser.add_argument("source", type=Path, help="Внутренний файл UT_PO_STATUS_<дата>.xlsx")
    parser.add_argument("--previous", type=Path, required=True, help="Прошлая клиентская книга")
    parser.add_argument("--template", type=Path, required=True, help="Шаблон с форматированием и списками")
    parser.add_argument("--output", type=Path, default=None)
    parser.add_argument("--history", type=Path, default=Path("state/po_status_history.json"))
    args = parser.parse_args()

    file_date = file_date_from_name(args.source)
    stamp = file_date.strftime("%d.%m.%Y")
    output = args.output or Path(f"Ready_UT_PO_STATUS_{stamp}.xlsx")
    changes_path = output.with_name(f"Ready_UT_PO_STATUS_{stamp}_changes.txt")

    source_rows = load_source(args.source)
    assign_batches(source_rows, file_date)
    previous = load_previous(args.previous)
    history = load_history(args.history)
    ordered, changes, dropped, notes = assemble(previous, source_rows, file_date, history)
    write_workbook(args.template, output, ordered, file_date)
    save_history(args.history, ordered, file_date, args.source.name)
    write_changes(changes_path, file_date, ordered, changes, dropped, notes)
    self_check(ordered, output, file_date)
    print(f"Готово: {output}")
    print(f"Отличия: {changes_path}")
    print(f"Изменений: {len(changes)}, убрано: {len(dropped)}")


if __name__ == "__main__":
    main()
