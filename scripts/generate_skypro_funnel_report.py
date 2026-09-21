#!/usr/bin/env python3
"""FASTAIR — Скайпро Техникс: воронка запросов (ТУЗ + Expendables).

Период: 01.06.2026 — текущая дата (по дате запроса / RFQ).
1 запрос = уникальная пара (P/N + дата столбца B).

Воронка: запрос → предложение → заказ.
Логика статусов/цен — как в отчётах UTair (ТУЗ / EXP).
"""

from __future__ import annotations

import argparse
import csv
import html
import math
import re
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl

PERIOD_START = date(2026, 6, 1)
DEFAULT_TUZ = Path("/tmp/skypro_analysis/ТУЗ 21.09.2026.xlsx")
DEFAULT_EXP = Path("/tmp/skypro_analysis/EXPENDABLES 21.09.2026.csv")
DEFAULT_TAZ = Path("/tmp/skypro_analysis/ТАЗ полный файл 28.08.2026.xlsx")
DEFAULT_OUT = Path("/workspace/output/skypro_technics_funnel.html")

GROUP_SHEETS = [
    "Группа A",
    "Группа B",
    "Группа C",
    "Группа 2 (old)",
    "Группа 3 (old)",
    "Группа 5 (old)",
]
OLD_SHEETS = {"Группа 2 (old)", "Группа 3 (old)", "Группа 5 (old)"}

NOT_FOUND_STATUSES = {
    "1. Не нашли",
    "2. Не будет проквотировано",
    "1. Вне компетенции",
    "1. Пропуск",
}
TAZ_EXCLUDED = {"5 CANCELLED", "7 REFUND"}
STATUS_PRIORITY = [
    "7. Клиент согласовал",
    "5. Есть интерес",
    "4. Квотация направлена клиенту",
    "3. Цена на деталь получена",
    "1. Проценка у поставщиков",
    "6. Клиент отказал",
    "8. Backup",
    "0. Начальный этап",
    "11. Внесено в ТАЗ",
]


def clean_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def normalize_pn(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "").strip().upper())


def is_skypro(customer: Any) -> bool:
    """Скайпро Техникс / Скайпро (не Глобал Скай, не Скай Партнер)."""
    if not customer:
        return False
    text = str(customer).strip().lower().replace("ё", "е")
    compact = re.sub(r"[\s\-_/]+", "", text)
    if "глобалскай" in compact or "globalsky" in compact:
        return False
    if "скайпартнер" in compact or "skypartner" in compact:
        return False
    if "скайпро" in compact or "skypro" in compact:
        return True
    if "скайротехник" in compact:
        return True
    return False


def to_dt(value: Any) -> datetime | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return None
    text = str(value).strip()
    for fmt in (
        "%d.%m.%Y %H:%M:%S",
        "%d.%m.%Y %H:%M",
        "%d.%m.%Y",
        "%d.%m.%y",
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d",
        "%m/%d/%Y",
    ):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def to_date(value: Any) -> date | None:
    dt = to_dt(value)
    return dt.date() if dt else None


def to_number(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
            return None
        return float(value)
    text = (
        str(value)
        .strip()
        .replace("\xa0", "")
        .replace(" ", "")
        .replace("$", "")
        .replace(",", ".")
    )
    if not text or set(text) <= {"-", "—"}:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def in_period(d: date | None, start: date, end: date) -> bool:
    return bool(d and start <= d <= end)


def pct(part: int, whole: int) -> float | None:
    if whole <= 0:
        return None
    return 100.0 * part / whole


def fmt_pct(v: float | None) -> str:
    if v is None:
        return "—"
    return f"{v:.0f}%"


def fmt_int(v: int) -> str:
    return f"{v:,}".replace(",", " ")


def sheet_columns(sheet: str) -> dict[str, int]:
    """1-based column map (openpyxl-style indices for values[i-1])."""
    if sheet in OLD_SHEETS:
        return {
            "status": 1,
            "request_dt": 2,
            "request_no": 3,
            "urgency": 4,
            "client": 5,
            "pn": 7,
            "alt_pn": 8,
            "description": 9,
            "qty": 10,
            "quote_dt": 12,
            "supplier_price": 16,
            "supplier": 17,
            "qty_avail": 21,
            "offered": 25,
            "sent_dt": 26,
            "accepted_dt": 27,
            "invoice": 28,
        }
    # Группа A/B/C: Sales ID / Purch ID at L/M, quote at O
    return {
        "status": 1,
        "request_dt": 2,
        "request_no": 3,
        "urgency": 4,
        "client": 5,
        "pn": 7,
        "alt_pn": 8,
        "description": 9,
        "qty": 10,
        "quote_dt": 15,
        "supplier_price": 19,
        "supplier": 20,
        "qty_avail": 24,
        "offered": 28,
        "sent_dt": 29,
        "accepted_dt": 30,
        "invoice": 31,
    }


def sheet_columns_shifted(sheet: str) -> dict[str, int]:
    if sheet in OLD_SHEETS:
        return sheet_columns(sheet)
    return {
        "status": 1,
        "request_dt": 4,
        "request_no": 5,
        "urgency": 3,
        "client": 7,
        "pn": 9,
        "alt_pn": 10,
        "description": 11,
        "qty": 12,
        "quote_dt": 15,
        "supplier_price": 19,
        "supplier": 20,
        "qty_avail": 24,
        "offered": 28,
        "sent_dt": 29,
        "accepted_dt": 30,
        "invoice": 31,
    }


def is_shifted_row(values: list) -> bool:
    if len(values) < 8:
        return False
    return is_skypro(values[6]) and not is_skypro(values[4])


def resolve_columns(sheet: str, values: list) -> dict[str, int]:
    if sheet in {"Группа A", "Группа B", "Группа C"} and is_shifted_row(values):
        return sheet_columns_shifted(sheet)
    return sheet_columns(sheet)


def cell(values: list, cols: dict[str, int], key: str):
    idx = cols[key] - 1
    return values[idx] if 0 <= idx < len(values) else None


@dataclass
class OfferRow:
    sheet: str
    status: str
    request_dt: datetime | None
    request_no: str
    urgency: str | None
    client: str
    pn: str
    description: str | None
    qty: float | None
    supplier_price: float | None
    supplier: str | None
    offered: float | None
    sent_dt: datetime | None
    invoice: str | None


@dataclass
class RequestAgg:
    pn: str
    request_date: date
    source: str  # tuz | exp
    offers: list[Any] = field(default_factory=list)
    description: str | None = None
    qty: float | None = None
    statuses: list[str] = field(default_factory=list)
    expr: str | None = None
    has_market: bool = False
    has_offer: bool = False
    has_order: bool = False
    order_refs: list[str] = field(default_factory=list)

    @property
    def primary_status(self) -> str:
        st = {s for s in self.statuses if s}
        for label in STATUS_PRIORITY:
            if label in st:
                return label
        return next(iter(st), "—")


@dataclass
class FunnelBlock:
    title: str
    subtitle: str
    requests: int
    offers: int
    orders: int
    found: int = 0
    by_month: list[tuple[str, int, int, int]] = field(default_factory=list)
    by_status: list[tuple[str, int]] = field(default_factory=list)
    sample_rows: list[dict[str, Any]] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)

    @property
    def offer_pct(self) -> float | None:
        return pct(self.offers, self.requests)

    @property
    def order_from_offer_pct(self) -> float | None:
        return pct(self.orders, self.offers)

    @property
    def order_from_request_pct(self) -> float | None:
        return pct(self.orders, self.requests)


def load_tuz_offers(path: Path) -> list[OfferRow]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out: list[OfferRow] = []
    for sheet in GROUP_SHEETS:
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        rows = ws.iter_rows(values_only=True)
        next(rows, None)
        for row in rows:
            values = list(row)
            cols = resolve_columns(sheet, values)
            client = clean_text(cell(values, cols, "client"))
            if not is_skypro(client):
                continue
            pn = normalize_pn(cell(values, cols, "pn"))
            if not pn:
                continue
            out.append(
                OfferRow(
                    sheet=sheet,
                    status=clean_text(cell(values, cols, "status")) or "",
                    request_dt=to_dt(cell(values, cols, "request_dt")),
                    request_no=clean_text(cell(values, cols, "request_no")) or "",
                    urgency=clean_text(cell(values, cols, "urgency")),
                    client=client or "",
                    pn=pn,
                    description=clean_text(cell(values, cols, "description")),
                    qty=to_number(cell(values, cols, "qty")),
                    supplier_price=to_number(cell(values, cols, "supplier_price")),
                    supplier=clean_text(cell(values, cols, "supplier")),
                    offered=to_number(cell(values, cols, "offered")),
                    sent_dt=to_dt(cell(values, cols, "sent_dt")),
                    invoice=clean_text(cell(values, cols, "invoice")),
                )
            )
    wb.close()
    return out


def aggregate_tuz(offers: list[OfferRow], start: date, end: date) -> list[RequestAgg]:
    grouped: dict[tuple[str, date], RequestAgg] = {}
    for offer in offers:
        req_date = offer.request_dt.date() if offer.request_dt else None
        if not in_period(req_date, start, end):
            continue
        key = (offer.pn, req_date)  # type: ignore[arg-type]
        agg = grouped.get(key)
        if not agg:
            agg = RequestAgg(
                pn=offer.pn,
                request_date=req_date,  # type: ignore[arg-type]
                source="tuz",
                description=offer.description,
                qty=offer.qty,
            )
            grouped[key] = agg
        agg.offers.append(offer)
        if offer.status:
            agg.statuses.append(offer.status)
        if offer.description and not agg.description:
            agg.description = offer.description
        if offer.qty is not None and agg.qty is None:
            agg.qty = offer.qty

        if (
            offer.status not in NOT_FOUND_STATUSES
            and offer.status != "8. Backup"
            and offer.supplier_price is not None
        ):
            agg.has_market = True

        # Предложение клиенту (UTair: sent / offered / статусы после квотации)
        if (
            offer.sent_dt
            or offer.offered is not None
            or offer.status.startswith(("4.", "5.", "6.", "7."))
        ):
            agg.has_offer = True

        if offer.status == "7. Клиент согласовал" or offer.invoice:
            agg.has_order = True
            if offer.invoice:
                agg.order_refs.append(offer.invoice)

    return sorted(grouped.values(), key=lambda r: (r.request_date, r.pn))


def load_exp_rows(path: Path) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    with path.open(encoding="utf-8-sig", newline="") as fh:
        reader = csv.reader(fh)
        next(reader, None)
        for raw in reader:
            if len(raw) < 43:
                continue
            client = clean_text(raw[3])
            if not is_skypro(client):
                continue
            pn = normalize_pn(raw[6])
            req_date = to_date(raw[1])
            if not pn or not req_date:
                continue
            market = to_number(raw[9])
            ddp_routine = to_number(raw[29] if len(raw) > 29 else None)
            ddp_critical = to_number(raw[38] if len(raw) > 38 else None)
            ddp = None
            for cand in (ddp_routine, ddp_critical):
                if cand is not None and cand > 0:
                    ddp = cand
                    break
            rows.append(
                {
                    "client": client,
                    "pn": pn,
                    "request_date": req_date,
                    "description": clean_text(raw[7]),
                    "qty": to_number(raw[8]),
                    "market": market,
                    "supplier": clean_text(raw[10]),
                    "ddp": ddp,
                    "expr": clean_text(raw[42]),
                    "remarks": clean_text(raw[4]),
                }
            )
    return rows


def aggregate_exp(rows: list[dict[str, Any]], start: date, end: date) -> list[RequestAgg]:
    grouped: dict[tuple[str, date], RequestAgg] = {}
    for row in rows:
        if not in_period(row["request_date"], start, end):
            continue
        key = (row["pn"], row["request_date"])
        agg = grouped.get(key)
        if not agg:
            agg = RequestAgg(
                pn=row["pn"],
                request_date=row["request_date"],
                source="exp",
                description=row["description"],
                qty=row["qty"],
                expr=row["expr"],
            )
            grouped[key] = agg
        agg.offers.append(row)
        if row["expr"] and not agg.expr:
            agg.expr = row["expr"]
        if row["market"] is not None:
            agg.has_market = True
        # UTair EXP: предложение = DDP > 0 (иначе market price как мягкий сигнал оффера)
        if row["ddp"] is not None or (row["market"] is not None and row["supplier"]):
            agg.has_offer = True
    return sorted(grouped.values(), key=lambda r: (r.request_date, r.pn))


@dataclass
class TazLine:
    invoice: str
    customer: str
    pn: str
    category: str
    work_date: date
    status: str
    sale_total: float | None


def load_taz_skypro(path: Path, start: date, end: date) -> list[TazLine]:
    if not path.exists():
        return []
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out: list[TazLine] = []
    for sheet in ("ORDERS", "PRESALE"):
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        rows = ws.iter_rows(values_only=True)
        next(rows, None)
        for row in rows:
            if not row or len(row) < 17:
                continue
            customer = clean_text(row[6])
            if not is_skypro(customer):
                continue
            status = (clean_text(row[4]) or "").upper()
            if status in TAZ_EXCLUDED:
                continue
            pn = normalize_pn(row[10])
            work = to_date(row[16])
            if not pn or not in_period(work, start, end):
                continue
            category = (clean_text(row[15]) or "").upper()
            invoice = clean_text(row[0]) or ""
            sale = to_number(row[33] if len(row) > 33 else None)
            out.append(
                TazLine(
                    invoice=invoice,
                    customer=customer or "",
                    pn=pn,
                    category=category,
                    work_date=work,  # type: ignore[arg-type]
                    status=status,
                    sale_total=sale,
                )
            )
    wb.close()
    return out


def match_exp_orders(requests: list[RequestAgg], taz: list[TazLine]) -> None:
    """Заказ EXP = тот же P/N в ТАЗ EXPENDABLE, дата заказа ≥ даты RFQ."""
    by_pn: dict[str, list[TazLine]] = defaultdict(list)
    for line in taz:
        if line.category == "EXPENDABLE":
            by_pn[line.pn].append(line)
    for req in requests:
        hits = [
            line
            for line in by_pn.get(req.pn, [])
            if line.work_date >= req.request_date
        ]
        if hits:
            req.has_order = True
            req.order_refs = sorted({h.invoice for h in hits if h.invoice})


def build_funnel(
    title: str,
    subtitle: str,
    requests: list[RequestAgg],
    *,
    notes: list[str] | None = None,
    include_status: bool = False,
) -> FunnelBlock:
    req_n = len(requests)
    offer_n = sum(1 for r in requests if r.has_offer)
    order_n = sum(1 for r in requests if r.has_order)
    found_n = sum(1 for r in requests if r.has_market)

    month_map: dict[str, list[int]] = defaultdict(lambda: [0, 0, 0])
    for r in requests:
        key = r.request_date.strftime("%Y-%m")
        month_map[key][0] += 1
        if r.has_offer:
            month_map[key][1] += 1
        if r.has_order:
            month_map[key][2] += 1
    by_month = [(m, *month_map[m]) for m in sorted(month_map)]

    by_status: list[tuple[str, int]] = []
    if include_status:
        c = Counter(r.primary_status for r in requests)
        by_status = sorted(c.items(), key=lambda x: (-x[1], x[0]))

    samples = []
    for r in requests[:80]:
        samples.append(
            {
                "date": r.request_date.isoformat(),
                "pn": r.pn,
                "description": r.description or "—",
                "qty": r.qty if r.qty is not None else "—",
                "status": r.primary_status if include_status else ("да" if r.has_offer else "нет"),
                "offer": "да" if r.has_offer else "нет",
                "order": "да" if r.has_order else "нет",
                "refs": ", ".join(r.order_refs[:3]) if r.order_refs else "—",
                "extra": r.expr or "",
            }
        )

    return FunnelBlock(
        title=title,
        subtitle=subtitle,
        requests=req_n,
        offers=offer_n,
        orders=order_n,
        found=found_n,
        by_month=by_month,
        by_status=by_status,
        sample_rows=samples,
        notes=notes or [],
    )


def render_funnel_kpis(block: FunnelBlock) -> str:
    return f"""
    <div class="kpis three">
      <div class="highlight">
        <div class="label">Запросы</div>
        <div class="value">{fmt_int(block.requests)}</div>
        <div class="muted">уник. P/N + дата B</div>
      </div>
      <div>
        <div class="label">Предложения</div>
        <div class="value">{fmt_int(block.offers)}</div>
        <div class="muted">{fmt_pct(block.offer_pct)} от запросов</div>
      </div>
      <div>
        <div class="label">Заказы</div>
        <div class="value">{fmt_int(block.orders)}</div>
        <div class="muted">{fmt_pct(block.order_from_offer_pct)} от предложений · {fmt_pct(block.order_from_request_pct)} от запросов</div>
      </div>
    </div>
    <div class="funnel-bar">
      <div class="stage s1"><span>Запрос</span><b>{fmt_int(block.requests)}</b></div>
      <div class="arrow">→ {fmt_pct(block.offer_pct)}</div>
      <div class="stage s2"><span>Предложение</span><b>{fmt_int(block.offers)}</b></div>
      <div class="arrow">→ {fmt_pct(block.order_from_offer_pct)}</div>
      <div class="stage s3"><span>Заказ</span><b>{fmt_int(block.orders)}</b></div>
    </div>
    """


def render_month_table(block: FunnelBlock, table_id: str) -> str:
    rows = []
    for month, reqs, offers, orders in block.by_month:
        rows.append(
            "<tr>"
            f"<td>{html.escape(month)}</td>"
            f"<td class='num'>{reqs}</td>"
            f"<td class='num'>{offers}</td>"
            f"<td class='num'>{fmt_pct(pct(offers, reqs))}</td>"
            f"<td class='num'>{orders}</td>"
            f"<td class='num'>{fmt_pct(pct(orders, offers))}</td>"
            "</tr>"
        )
    body = "\n".join(rows) or "<tr><td colspan='6' class='empty'>Нет данных</td></tr>"
    return f"""
    <div class="table-scroll">
    <table data-sortable id="{html.escape(table_id)}">
      <thead><tr>
        <th>Месяц <span class="arrow">↕</span></th>
        <th>Запросы <span class="arrow">↕</span></th>
        <th>Предложения <span class="arrow">↕</span></th>
        <th>Запрос→Предл. <span class="arrow">↕</span></th>
        <th>Заказы <span class="arrow">↕</span></th>
        <th>Предл.→Заказ <span class="arrow">↕</span></th>
      </tr></thead>
      <tbody>{body}</tbody>
    </table>
    </div>
    """


def render_status_table(block: FunnelBlock, table_id: str) -> str:
    if not block.by_status:
        return ""
    rows = []
    for status, count in block.by_status:
        rows.append(
            "<tr>"
            f"<td>{html.escape(status)}</td>"
            f"<td class='num'>{count}</td>"
            f"<td class='num'>{fmt_pct(pct(count, block.requests))}</td>"
            "</tr>"
        )
    return f"""
    <details class="subblock" open>
      <summary>Статусы запросов (ТУЗ)</summary>
      <div class="block-body">
        <div class="table-scroll">
        <table data-sortable id="{html.escape(table_id)}">
          <thead><tr>
            <th>Статус <span class="arrow">↕</span></th>
            <th>Запросов <span class="arrow">↕</span></th>
            <th>Доля <span class="arrow">↕</span></th>
          </tr></thead>
          <tbody>{''.join(rows)}</tbody>
        </table>
        </div>
      </div>
    </details>
    """


def render_sample_table(block: FunnelBlock, table_id: str, *, tuz: bool) -> str:
    rows = []
    for item in block.sample_rows:
        rows.append(
            "<tr>"
            f"<td>{html.escape(item['date'])}</td>"
            f"<td>{html.escape(item['pn'])}</td>"
            f"<td>{html.escape(str(item['description']))}</td>"
            f"<td class='num'>{html.escape(str(item['qty']))}</td>"
            + (f"<td>{html.escape(str(item['status']))}</td>" if tuz else "")
            + f"<td>{html.escape(item['offer'])}</td>"
            f"<td>{html.escape(item['order'])}</td>"
            f"<td>{html.escape(str(item['refs']))}</td>"
            "</tr>"
        )
    status_col = "<th>Статус <span class='arrow'>↕</span></th>" if tuz else ""
    body = "\n".join(rows) or "<tr><td colspan='8' class='empty'>Нет строк</td></tr>"
    return f"""
    <details class="subblock">
      <summary>Примеры запросов (первые {len(block.sample_rows)})</summary>
      <div class="block-body">
        <div class="table-scroll">
        <table data-sortable id="{html.escape(table_id)}">
          <thead><tr>
            <th>Дата B <span class="arrow">↕</span></th>
            <th>P/N <span class="arrow">↕</span></th>
            <th>Описание <span class="arrow">↕</span></th>
            <th>Qty <span class="arrow">↕</span></th>
            {status_col}
            <th>Предложение <span class="arrow">↕</span></th>
            <th>Заказ <span class="arrow">↕</span></th>
            <th>Счёт / ссылка <span class="arrow">↕</span></th>
          </tr></thead>
          <tbody>{body}</tbody>
        </table>
        </div>
      </div>
    </details>
    """


def render_section(block: FunnelBlock, *, opened: bool, tuz: bool, idx: int) -> str:
    open_attr = " open" if opened else ""
    notes = ""
    if block.notes:
        notes = (
            "<div class='note-box'><strong>Примечания</strong><ul>"
            + "".join(f"<li>{html.escape(n)}</li>" for n in block.notes)
            + "</ul></div>"
        )
    extra = ""
    if tuz and block.found:
        extra = f"<p class='hint'>Найдено на рынке (цена поставщика, без Backup / not-found): <b>{fmt_int(block.found)}</b> из {fmt_int(block.requests)} ({fmt_pct(pct(block.found, block.requests))}).</p>"
    return f"""
<details class="period"{open_attr}>
  <summary>
    <span class="period-title">{html.escape(block.title)}</span>
    <span class="period-meta">{html.escape(block.subtitle)}</span>
    <span class="period-meta">{fmt_int(block.requests)} запросов · {fmt_int(block.offers)} предложений · {fmt_int(block.orders)} заказов</span>
  </summary>
  <div class="period-body">
    {render_funnel_kpis(block)}
    {extra}
    <details class="block" open>
      <summary>1. Воронка по месяцам</summary>
      <div class="block-body">{render_month_table(block, f"m{idx}")}</div>
    </details>
    {render_status_table(block, f"s{idx}") if tuz else ""}
    {render_sample_table(block, f"r{idx}", tuz=tuz)}
    {notes}
  </div>
</details>
"""


def render_html(
    *,
    tuz_block: FunnelBlock,
    exp_block: FunnelBlock,
    period_label: str,
    sources: dict[str, str],
) -> str:
    sections = render_section(tuz_block, opened=True, tuz=True, idx=1) + render_section(
        exp_block, opened=True, tuz=False, idx=2
    )
    return f"""<!DOCTYPE html>
<html lang="ru">
<head>
<meta charset="utf-8"/>
<meta name="viewport" content="width=device-width, initial-scale=1"/>
<title>FASTAIR — Скайпро Техникс · воронка</title>
<style>
:root {{
  --bg:#e7ecef; --card:#fff; --ink:#202020; --muted:#5a5a5a;
  --navy:#022f40; --cyan:#d5fbff; --line:#c4c4c4; --zebra:#f5fafb;
}}
* {{ box-sizing:border-box; }}
body {{
  margin:0; font-family: Arial, Calibri, Helvetica, sans-serif;
  color:var(--ink); background: linear-gradient(180deg, #022f40 0 140px, var(--bg) 140px);
  line-height:1.45;
}}
.wrap {{ max-width:1100px; margin:0 auto; padding:28px 20px 64px; }}
.brand {{
  background:var(--cyan); color:var(--navy); display:inline-block;
  padding:6px 10px; font-weight:700; margin-bottom:10px;
  font-size:13px; letter-spacing:.14em; text-transform:uppercase;
}}
h1 {{ color:#fff; font-size:28px; margin:0 0 6px; font-weight:700; }}
.sub {{ color:#d5fbff; margin:0 0 18px; font-size:14px; }}
details.period {{
  background:var(--card); border:1px solid var(--line); border-radius:8px;
  margin-bottom:14px; overflow:hidden;
}}
details.period > summary {{
  cursor:pointer; list-style:none; padding:14px 16px;
  background:#f3f8fa; border-bottom:1px solid transparent;
  display:flex; flex-wrap:wrap; gap:8px 16px; align-items:baseline;
}}
details.period[open] > summary {{ border-bottom-color:var(--line); }}
details.period > summary::-webkit-details-marker {{ display:none; }}
.period-title {{ font-size:17px; font-weight:700; color:var(--navy); }}
.period-meta {{ font-size:13px; color:var(--muted); }}
.period-body {{ padding:12px 16px 16px; }}
details.block, details.subblock {{
  border:1px solid #d7e2e6; border-radius:8px; margin:10px 0; background:#fff;
}}
details.block > summary, details.subblock > summary {{
  cursor:pointer; padding:10px 12px; font-weight:700; color:var(--navy);
  list-style:none; background:#f7fbfc;
}}
details.block > summary::-webkit-details-marker,
details.subblock > summary::-webkit-details-marker {{ display:none; }}
details.block > summary::before,
details.subblock > summary::before {{ content:"▸ "; color:#7a8a90; }}
details.block[open] > summary::before,
details.subblock[open] > summary::before {{ content:"▾ "; }}
.block-body {{ padding:12px; }}
.kpis {{ display:grid; grid-template-columns:repeat(4,minmax(0,1fr)); gap:10px; margin-bottom:14px; }}
.kpis.three {{ grid-template-columns:repeat(3,minmax(0,1fr)); }}
.kpis > div {{ background:#f7fbfc; border:1px solid #d7e2e6; border-radius:8px; padding:12px 14px; }}
.kpis > div.highlight {{ background:#e8f6f8; border-color:#9ed7e0; }}
.label {{ font-size:11px; text-transform:uppercase; letter-spacing:.04em; color:#757575; }}
.value {{ font-size:22px; font-weight:700; margin-top:4px; color:var(--navy); font-variant-numeric:tabular-nums; }}
.muted {{ color:var(--muted); font-size:12px; margin-top:4px; }}
.funnel-bar {{
  display:flex; flex-wrap:wrap; align-items:center; gap:8px; margin:8px 0 16px;
}}
.funnel-bar .stage {{
  background:var(--navy); color:#fff; border-radius:8px; padding:10px 14px; min-width:120px;
}}
.funnel-bar .stage.s2 {{ background:#03425a; }}
.funnel-bar .stage.s3 {{ background:#0a5c4c; }}
.funnel-bar .stage span {{ display:block; font-size:11px; opacity:.85; text-transform:uppercase; letter-spacing:.04em; }}
.funnel-bar .stage b {{ font-size:20px; }}
.funnel-bar .arrow {{ color:var(--muted); font-size:13px; font-weight:700; }}
.table-scroll {{ overflow-x:auto; }}
table {{ width:100%; border-collapse:collapse; font-size:14px; }}
th, td {{ border-bottom:1px solid #e8e8e8; padding:9px 8px; text-align:left; vertical-align:middle; }}
th {{
  font-size:12px; text-transform:uppercase; letter-spacing:.03em; color:#fff;
  background:var(--navy); cursor:pointer; user-select:none; white-space:nowrap;
}}
th:hover {{ background:#03425a; }}
th .arrow {{ opacity:.45; margin-left:6px; font-size:11px; }}
th.sorted .arrow {{ opacity:1; }}
td.num {{ text-align:right; font-variant-numeric:tabular-nums; white-space:nowrap; }}
td.empty, .empty {{ color:#9a9a9a; }}
tbody tr:nth-child(even) {{ background:var(--zebra); }}
tbody tr:hover {{ background:#eef7f9; }}
.hint {{ margin-top:12px; color:var(--muted); font-size:12px; }}
.rules {{
  background:#e8f6f8; border:1px solid #9ed7e0; border-radius:8px;
  padding:12px 14px; margin-bottom:16px; font-size:13px; color:var(--navy);
}}
.note-box {{
  background:#fff8f2; border:1px solid #f0c7a8; border-radius:8px;
  padding:12px 14px; margin-top:12px; font-size:13px; color:#5a3a22;
}}
.note-box ul {{ margin:8px 0 0; padding-left:18px; }}
.note-box li {{ margin:4px 0; }}
@media (max-width:800px) {{
  .kpis, .kpis.three {{ grid-template-columns:repeat(2,minmax(0,1fr)); }}
}}
</style>
</head>
<body>
<div class="wrap">
  <div class="brand">FASTAIR</div>
  <h1>Скайпро Техникс — воронка</h1>
  <div class="sub">Запрос → предложение → заказ · период {html.escape(period_label)} · два потока: ТУЗ и Expendables</div>
  <div class="rules">
    <strong>Клиент:</strong> Скайпро Техникс / Скайпро (без «Глобал Скай» и «Скай Партнер»).
    <br/>
    <strong>1 запрос</strong> = уникальная пара <strong>P/N + дата из столбца B</strong>
    (в EXP — Date of RFQ).
    <br/>
    <strong>ТУЗ · предложение:</strong> Offered / Sent to client / статусы 4–7
    (как в анализе UTair). <strong>Заказ:</strong> статус «7. Клиент согласовал» или заполнен Invoice.
    <br/>
    <strong>Expendables · предложение:</strong> DDP-цена клиенту &gt; 0 (или market price + supplier).
    <strong>Заказ:</strong> тот же P/N в ТАЗ Category=EXPENDABLE, дата заказа ≥ даты RFQ
    (логика UTair EXP).
  </div>
  {sections}
  <p class="hint">
    Источники: ТУЗ — {html.escape(sources.get('tuz','—'))};
    Expendables — {html.escape(sources.get('exp','—'))};
    ТАЗ (для заказов EXP) — {html.escape(sources.get('taz','—'))}.
    Скачивайте HTML напрямую (ZIP+HTML Windows часто помечает ложно).
  </p>
</div>
<script>
document.querySelectorAll('table[data-sortable]').forEach((table) => {{
  const headers = table.querySelectorAll('th');
  headers.forEach((th, idx) => {{
    th.addEventListener('click', () => {{
      const tbody = table.tBodies[0];
      if (!tbody) return;
      const rows = Array.from(tbody.querySelectorAll('tr'));
      const asc = th.dataset.asc !== '1';
      headers.forEach(h => {{ h.classList.remove('sorted'); h.dataset.asc = ''; }});
      th.classList.add('sorted');
      th.dataset.asc = asc ? '1' : '0';
      const parse = (td) => {{
        const t = (td?.textContent || '').trim().replace(/\\s/g, '').replace('%','');
        const n = Number(t.replace(',', '.'));
        return Number.isFinite(n) && t !== '' && t !== '—' ? n : t;
      }};
      rows.sort((a, b) => {{
        const av = parse(a.children[idx]);
        const bv = parse(b.children[idx]);
        if (typeof av === 'number' && typeof bv === 'number') return asc ? av - bv : bv - av;
        return asc ? String(av).localeCompare(String(bv), 'ru') : String(bv).localeCompare(String(av), 'ru');
      }});
      rows.forEach(r => tbody.appendChild(r));
    }});
  }});
}});
</script>
</body>
</html>
"""


def main() -> None:
    parser = argparse.ArgumentParser(description="Skypro Technics funnel HTML report")
    parser.add_argument("--tuz", type=Path, default=DEFAULT_TUZ)
    parser.add_argument("--exp", type=Path, default=DEFAULT_EXP)
    parser.add_argument("--taz", type=Path, default=DEFAULT_TAZ)
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT)
    parser.add_argument("--end", type=str, default="", help="YYYY-MM-DD (default: today)")
    args = parser.parse_args()

    end = date.fromisoformat(args.end) if args.end else date.today()
    start = PERIOD_START
    period_label = f"{start.strftime('%d.%m.%Y')} – {end.strftime('%d.%m.%Y')}"

    tuz_offers = load_tuz_offers(args.tuz)
    tuz_reqs = aggregate_tuz(tuz_offers, start, end)
    tuz_block = build_funnel(
        "1. ТУЗ (rotables / компоненты)",
        f"листы групп · {args.tuz.name}",
        tuz_reqs,
        include_status=True,
        notes=[
            "Уникальность запроса: P/N + дата Request date&time (столбец B), не Request №.",
            "Несколько строк поставщиков по одному P/N+дате схлопываются в один запрос.",
        ],
    )

    exp_rows = load_exp_rows(args.exp)
    exp_reqs = aggregate_exp(exp_rows, start, end)
    taz_lines = load_taz_skypro(args.taz, start, end)
    match_exp_orders(exp_reqs, taz_lines)
    exp_rotable = sum(1 for t in taz_lines if t.category == "ROTABLE")
    exp_expendable = sum(1 for t in taz_lines if t.category == "EXPENDABLE")
    exp_block = build_funnel(
        "2. Expendables",
        f"CSV · {args.exp.name}",
        exp_reqs,
        notes=[
            "Предложение: DDP-цена клиенту > 0, иначе market price + supplier (как в UTair EXP).",
            f"Заказы сверстаны с ТАЗ EXPENDABLE ({args.taz.name}): строк Скайпро в периоде — "
            f"EXPENDABLE {exp_expendable}, ROTABLE {exp_rotable} (ROTABLE в эту воронку не входят).",
            "Срез ТАЗ от 28.08.2026: заказы сентября после этой даты могут быть неполными.",
        ],
    )

    args.out.parent.mkdir(parents=True, exist_ok=True)
    html_text = render_html(
        tuz_block=tuz_block,
        exp_block=exp_block,
        period_label=period_label,
        sources={
            "tuz": args.tuz.name,
            "exp": args.exp.name,
            "taz": args.taz.name if args.taz.exists() else "нет файла",
        },
    )
    args.out.write_text(html_text, encoding="utf-8")
    print(f"Generated {args.out}")
    print(
        {
            "period": period_label,
            "tuz_requests": tuz_block.requests,
            "tuz_offers": tuz_block.offers,
            "tuz_orders": tuz_block.orders,
            "exp_requests": exp_block.requests,
            "exp_offers": exp_block.offers,
            "exp_orders": exp_block.orders,
        }
    )


if __name__ == "__main__":
    main()
